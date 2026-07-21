import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Read CSV
rows = []
with open("(26.07.20) 치기공 소프트웨어 가격 조사.csv", encoding="utf-8-sig", newline="") as f:
    reader = csv.DictReader(f)
    for row in reader:
        if any(row.values()):
            rows.append(row)

headers = ["제품명", "회사", "카테고리", "지역", "가격", "가격유형", "비고"]

# Define segments
segments = [
    ("국내-① 유사포지션",    lambda r: r["지역"] == "국내"  and r["카테고리"] == "1_유사포지션"),
    ("국내-② 규모있는SW",    lambda r: r["지역"] == "국내"  and r["카테고리"] == "2_규모있는소프트웨어"),
    ("국내-③ 전체",          lambda r: r["지역"] == "국내"),
    ("글로벌-① 유사포지션",  lambda r: r["지역"] == "글로벌" and r["카테고리"] == "1_유사포지션"),
    ("글로벌-② 규모있는SW",  lambda r: r["지역"] == "글로벌" and r["카테고리"] == "2_규모있는소프트웨어"),
    ("글로벌-③ 전체",        lambda r: r["지역"] == "글로벌"),
    ("전체",                  lambda r: True),
]

# Styles
header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(bold=True, color="FFFFFF", size=11)
alt_fill    = PatternFill("solid", fgColor="D6E4F0")
border_side = Side(style="thin", color="BFBFBF")
thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

def style_sheet(ws, data_rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for i, row in enumerate(data_rows, start=2):
        ws.append([row.get(h, "") for h in headers])
        fill = alt_fill if i % 2 == 0 else PatternFill()
        for cell in ws[i]:
            cell.fill = fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = thin_border

    # Column widths
    col_widths = [22, 22, 22, 10, 14, 14, 50]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

wb = openpyxl.Workbook()
wb.remove(wb.active)  # remove default sheet

for sheet_name, condition in segments:
    ws = wb.create_sheet(title=sheet_name)
    filtered = [r for r in rows if condition(r)]
    style_sheet(ws, filtered)

out_path = "(26.07.20) 치기공 소프트웨어 가격 조사.xlsx"
wb.save(out_path)
print(f"저장 완료: {out_path}")
