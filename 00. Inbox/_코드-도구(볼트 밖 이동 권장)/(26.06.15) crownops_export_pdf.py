from __future__ import annotations

import csv
import io
import os
import re
import shutil
import tempfile
import zipfile
import sys
from datetime import datetime
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator

from bs4 import BeautifulSoup
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from PIL import Image as PILImage, ImageOps
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from xml.sax.saxutils import escape


DATE_STR = "2026-06-15"
OWNER = "엄주원"
PREFIX = f"2026 GCS 8기_{OWNER}"
ZIP_PATH = Path(r"C:\Users\rweum\OneDrive\바탕 화면\2차 팀 크라운옵스.zip")
INNER_NAME = "ExportBlock-1cb2cd67-f135-475a-a9ea-d26d55baac73-Part-1.zip"
TARGET_PREFIX = "개인 페이지 & 공유된 페이지/자료실"
OUT_DIR = Path.home() / "Downloads" / "crownops_pdf_out_notiondate"
INDEX_CSV = OUT_DIR / f"{PREFIX}_index.csv"
INDEX_MD = OUT_DIR / f"{PREFIX}_index.md"
FONT_DIR = Path(r"C:\Windows\Fonts")
CHROME_PATH = Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe")


@dataclass(frozen=True)
class SourceFile:
    rel: str
    path: Path
    ext: str


CREATED_RE = re.compile(r"^생성 일시:\s*(.+)$", re.MULTILINE)
CSV_DATE_RE = re.compile(r"^([^,]+),([^,]*),([^,]*),(.+)$")


def register_fonts() -> str:
    normal = FONT_DIR / "malgun.ttf"
    bold = FONT_DIR / "malgunbd.ttf"
    if not normal.exists() or not bold.exists():
        raise FileNotFoundError("Malgun Gothic font files not found in C:\\Windows\\Fonts")
    pdfmetrics.registerFont(TTFont("Malgun", str(normal)))
    pdfmetrics.registerFont(TTFont("Malgun-Bold", str(bold)))
    pdfmetrics.registerFontFamily("Malgun", normal="Malgun", bold="Malgun-Bold")
    return "Malgun"


class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._states = []

    def showPage(self):
        self._states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        page_count = len(self._states)
        for state in self._states:
            self.__dict__.update(state)
            self.draw_footer(page_count)
            super().showPage()
        super().save()

    def draw_footer(self, page_count: int):
        self.setFont("Malgun", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        self.drawRightString(A4[0] - 16 * mm, 9 * mm, f"{self._pageNumber} / {page_count}")


def make_styles(font_name: str) -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "title",
            parent=base["Title"],
            fontName="Malgun-Bold",
            fontSize=20,
            leading=28,
            textColor=colors.HexColor("#0f172a"),
            wordWrap="CJK",
            spaceAfter=8,
        ),
        "meta": ParagraphStyle(
            "meta",
            parent=base["Normal"],
            fontName=font_name,
            fontSize=8.5,
            leading=13,
            textColor=colors.HexColor("#475569"),
            wordWrap="CJK",
            spaceAfter=4,
        ),
        "h1": ParagraphStyle(
            "h1",
            parent=base["Heading1"],
            fontName="Malgun-Bold",
            fontSize=16,
            leading=22,
            textColor=colors.HexColor("#0f172a"),
            wordWrap="CJK",
            spaceBefore=8,
            spaceAfter=6,
        ),
        "h2": ParagraphStyle(
            "h2",
            parent=base["Heading2"],
            fontName="Malgun-Bold",
            fontSize=13.5,
            leading=19,
            textColor=colors.HexColor("#111827"),
            wordWrap="CJK",
            spaceBefore=7,
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "body",
            parent=base["BodyText"],
            fontName=font_name,
            fontSize=10.2,
            leading=15.5,
            textColor=colors.HexColor("#111827"),
            wordWrap="CJK",
            spaceAfter=3,
        ),
        "small": ParagraphStyle(
            "small",
            parent=base["BodyText"],
            fontName=font_name,
            fontSize=8.2,
            leading=12,
            textColor=colors.HexColor("#475569"),
            wordWrap="CJK",
        ),
        "code": ParagraphStyle(
            "code",
            parent=base["Code"],
            fontName=font_name,
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#0f172a"),
            wordWrap="CJK",
        ),
    }


def p(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(text).replace("\n", "<br/>"), style)


def sanitize_filename(text: str, limit: int = 64) -> str:
    text = re.sub(r"[\\/:*?\"<>|]", "_", text)
    text = re.sub(r"\s+", " ", text).strip(" .\t")
    text = re.sub(r"_+", "_", text)
    if not text:
        text = "문서"
    if len(text) > limit:
        text = text[:limit].rstrip(" ._")
    return text


def normalize_date_string(raw: str) -> str | None:
    raw = raw.strip()
    if not raw:
        return None
    raw = raw.replace("오전", "AM").replace("오후", "PM")
    raw = raw.replace("년 ", "-").replace("월 ", "-").replace("일 ", " ")
    raw = raw.replace("년", "-").replace("월", "-").replace("일", " ")
    raw = raw.replace("  ", " ")
    candidates = [
        "%Y-%m-%d %p %I:%M",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
    ]
    for fmt in candidates:
        try:
            dt = datetime.strptime(raw, fmt)
            return dt.strftime("%Y-%m-%d")
        except ValueError:
            continue
    m = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", raw)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    return None


def extract_created_date_from_text(text: str) -> str | None:
    m = CREATED_RE.search(text)
    if m:
        return normalize_date_string(m.group(1))
    return None


def extract_created_date_from_csv(path: Path) -> str | None:
    text = path.read_text(encoding="utf-8-sig", errors="ignore")
    lines = [line for line in text.splitlines() if line.strip()]
    if len(lines) < 2:
        return None
    header = [h.strip() for h in next(csv.reader([lines[0]]))]
    rows = list(csv.reader(lines[1:]))
    if not rows:
        return None
    # Prefer an explicit created date column if present.
    date_col = None
    for idx, name in enumerate(header):
        if name in {"생성 일시", "생성", "생성 1"}:
            date_col = idx
            break
    if date_col is None:
        return None
    for row in rows:
        if len(row) > date_col:
            dt = normalize_date_string(row[date_col])
            if dt:
                return dt
    return None


def normalized_rel(path: Path, root: Path) -> str:
    return path.relative_to(root).as_posix()


def extract_outer_inner_zip() -> Path:
    if not ZIP_PATH.exists():
        raise FileNotFoundError(ZIP_PATH)
    temp_root = Path(tempfile.mkdtemp(prefix="crownops_extract_"))
    with zipfile.ZipFile(ZIP_PATH) as outer:
        data = outer.read(INNER_NAME)
    with zipfile.ZipFile(io.BytesIO(data)) as inner:
        inner.extractall(temp_root)
    return temp_root


def iter_source_files(root: Path) -> list[SourceFile]:
    items: list[SourceFile] = []
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        rel = normalized_rel(path, root)
        if not rel.startswith(TARGET_PREFIX):
            continue
        items.append(SourceFile(rel=rel, path=path, ext=path.suffix.lower()))
    items.sort(key=lambda item: item.rel)
    return items


def markdownish_title(text: str) -> str | None:
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        m = re.match(r"^#{1,6}\s+(.+)$", stripped)
        if m:
            return m.group(1).strip()
        if len(stripped) <= 80 and not stripped.startswith(("```", "- ", "* ", ">")):
            return stripped
    return None


def html_title(text: str) -> str | None:
    soup = BeautifulSoup(text, "html.parser")
    if soup.title and soup.title.get_text(strip=True):
        return soup.title.get_text(" ", strip=True)
    for tag in soup.find_all(["h1", "h2", "h3"]):
        t = tag.get_text(" ", strip=True)
        if t:
            return t
    return None


def docx_title(doc: Document) -> str | None:
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style and para.style.name and para.style.name.lower().startswith("heading"):
            return text
        return text
    return None


def txt_title(text: str) -> str | None:
    for line in text.splitlines():
        stripped = line.strip()
        if stripped:
            return stripped[:80]
    return None


def derive_title(source: SourceFile) -> str:
    ext = source.ext
    if ext in {".md", ".txt"}:
        raw = source.path.read_text(encoding="utf-8", errors="ignore")
        title = markdownish_title(raw) if ext == ".md" else txt_title(raw)
    elif ext == ".html":
        raw = source.path.read_text(encoding="utf-8", errors="ignore")
        title = html_title(raw)
    elif ext == ".docx":
        doc = Document(str(source.path))
        title = docx_title(doc)
    elif ext == ".csv":
        title = source.path.stem
    else:
        title = source.path.stem

    if not title:
        title = source.path.stem

    title = re.sub(r"\s+", " ", title).strip()
    title = re.sub(r"\b[0-9a-f]{8,}\b", "", title, flags=re.IGNORECASE)
    title = re.sub(r"[_\-]{2,}", " ", title).strip()
    if not title:
        title = source.path.stem
    return sanitize_filename(title)


def derive_created_date(source: SourceFile) -> str | None:
    ext = source.ext
    if ext == ".md":
        raw = source.path.read_text(encoding="utf-8", errors="ignore")
        return extract_created_date_from_text(raw)
    if ext == ".csv":
        return extract_created_date_from_csv(source.path)
    return None


def build_source_index(sources: list[SourceFile]) -> dict[str, SourceFile]:
    return {source.rel: source for source in sources}


def find_parent_page_date(
    source: SourceFile,
    source_index: dict[str, SourceFile],
    date_map: dict[str, str | None],
    root: Path,
) -> str | None:
    current_dir = source.path.parent
    while current_dir != root and current_dir.parent != current_dir:
        page_prefix = current_dir.name
        parent_dir = current_dir.parent
        candidates = [
            s
            for s in source_index.values()
            if s.path.parent == parent_dir
            and s.ext in {".md", ".csv"}
            and s.path.name.startswith(page_prefix)
        ]
        for cand in candidates:
            cand_date = date_map.get(cand.rel)
            if cand_date:
                return cand_date
        current_dir = parent_dir
    return None


def make_output_name(title: str, used: dict[str, int]) -> str:
    base = f"{PREFIX}_{title}"
    base = sanitize_filename(base, limit=180)
    count = used.get(base, 0)
    used[base] = count + 1
    if count == 0:
        return f"{base}.pdf"
    return f"{base}_{count + 1}.pdf"


def escape_text(text: str) -> str:
    return escape(text).replace("\n", "<br/>")


def maybe_linkify(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text


def fit_image(path: Path, max_w: float, max_h: float) -> Image:
    with PILImage.open(path) as img:
        width, height = img.size
    scale = min(max_w / width, max_h / height, 1.0)
    return Image(str(path), width=width * scale, height=height * scale)


def image_page(path: Path, styles: dict[str, ParagraphStyle], caption: str | None = None) -> list:
    flow: list = [p(path.name, styles["h2"])]
    if caption:
        flow.append(p(caption, styles["small"]))
    flow.append(Spacer(1, 3))
    flow.append(fit_image(path, 168 * mm, 220 * mm))
    return flow


def markdown_flowables(text: str, source_dir: Path, styles: dict[str, ParagraphStyle]) -> list:
    flow: list = []
    in_code = False
    code_buf: list[str] = []
    bullet_buf: list[str] = []

    def flush_bullets():
        nonlocal bullet_buf
        for item in bullet_buf:
            flow.append(p("• " + item, styles["body"]))
        bullet_buf = []

    def flush_code():
        nonlocal code_buf
        if code_buf:
            flow.append(Preformatted("\n".join(code_buf), styles["code"]))
        code_buf = []

    img_pattern = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")
    lines = text.splitlines()
    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                flush_code()
                in_code = False
            else:
                flush_bullets()
                in_code = True
            continue

        if in_code:
            code_buf.append(line)
            continue

        if not stripped:
            flush_bullets()
            flow.append(Spacer(1, 3))
            continue

        m = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if m:
            flush_bullets()
            flush_code()
            level = len(m.group(1))
            title = maybe_linkify(m.group(2).strip())
            flow.append(p(title, styles["h1"] if level <= 2 else styles["h2"]))
            continue

        m = re.match(r"^[-*+]\s+(.+)$", stripped)
        if m:
            flush_code()
            bullet_buf.append(m.group(1).strip())
            continue

        if re.match(r"^\d+[.)]\s+(.+)$", stripped):
            flush_bullets()
            flush_code()
            flow.append(p(maybe_linkify(stripped), styles["body"]))
            continue

        flush_bullets()
        flush_code()

        img_matches = list(img_pattern.finditer(line))
        if img_matches:
            before = img_pattern.sub("", line).strip()
            if before:
                flow.append(p(maybe_linkify(before), styles["body"]))
            for match in img_matches:
                img_path = (source_dir / match.group(1)).resolve()
                if img_path.exists() and img_path.is_file():
                    flow.extend(image_page(img_path, styles, caption=match.group(1)))
                else:
                    flow.append(p(f"[missing image] {match.group(1)}", styles["small"]))
            continue

        cleaned = maybe_linkify(line)
        cleaned = cleaned.replace("**", "").replace("__", "").replace("`", "")
        cleaned = re.sub(r"^>\s*", "", cleaned)
        flow.append(p(cleaned, styles["body"]))

    flush_bullets()
    flush_code()
    return flow


def html_flowables(text: str, source_dir: Path, styles: dict[str, ParagraphStyle]) -> list:
    soup = BeautifulSoup(text, "html.parser")
    flow: list = []

    for node in soup.find_all(["h1", "h2", "h3", "p", "li", "blockquote", "pre"]):
        txt = node.get_text(" ", strip=True)
        if not txt:
            continue
        if node.name == "h1":
            flow.append(p(txt, styles["h1"]))
        elif node.name == "h2" or node.name == "h3":
            flow.append(p(txt, styles["h2"]))
        elif node.name == "li":
            flow.append(p("• " + txt, styles["body"]))
        elif node.name == "blockquote":
            flow.append(p("> " + txt, styles["body"]))
        elif node.name == "pre":
            flow.append(Preformatted(txt, styles["code"]))
        else:
            flow.append(p(txt, styles["body"]))

    for img in soup.find_all("img"):
        src = img.get("src")
        if not src:
            continue
        img_path = (source_dir / src).resolve()
        if img_path.exists() and img_path.is_file():
            flow.extend(image_page(img_path, styles, caption=src))
    return flow


def csv_flowables(path: Path, styles: dict[str, ParagraphStyle]) -> list:
    with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        return [p("(빈 CSV)", styles["body"])]
    data = [[p(cell, styles["small"]) for cell in row] for row in rows]
    max_cols = max(len(row) for row in rows)
    for row in data:
        if len(row) < max_cols:
            row.extend([p("", styles["small"]) for _ in range(max_cols - len(row))])
    widths = [170 * mm / max_cols] * max_cols
    return [
        Table(
            data,
            colWidths=widths,
            repeatRows=1,
            style=TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eff6ff")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                    ("FONTNAME", (0, 0), (-1, 0), "Malgun-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            ),
        )
    ]


def iter_block_items(doc: DocxDocument) -> Iterator[DocxParagraph | DocxTable]:
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield DocxParagraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield DocxTable(child, doc)


def docx_table_to_flowable(table: DocxTable, styles: dict[str, ParagraphStyle]) -> Table:
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            cell_text = "\n".join(p.text for p in cell.paragraphs if p.text.strip())
            cells.append(p(cell_text or " ", styles["small"]))
        rows.append(cells)
    max_cols = max((len(r) for r in rows), default=1)
    for row in rows:
        if len(row) < max_cols:
            row.extend([p("", styles["small"]) for _ in range(max_cols - len(row))])
    widths = [170 * mm / max_cols] * max_cols
    return Table(
        rows,
        colWidths=widths,
        repeatRows=1,
        style=TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eff6ff")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                ("FONTNAME", (0, 0), (-1, 0), "Malgun-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        ),
    )


def docx_flowables(path: Path, styles: dict[str, ParagraphStyle]) -> list:
    doc = Document(str(path))
    flow: list = []
    for block in iter_block_items(doc):
        if isinstance(block, DocxParagraph):
            txt = block.text.strip()
            if not txt:
                flow.append(Spacer(1, 2))
                continue
            style_name = (block.style.name or "").lower() if block.style else ""
            if style_name.startswith("heading 1"):
                flow.append(p(txt, styles["h1"]))
            elif style_name.startswith("heading 2") or style_name.startswith("heading 3"):
                flow.append(p(txt, styles["h2"]))
            else:
                flow.append(p(txt, styles["body"]))
        elif isinstance(block, DocxTable):
            flow.append(docx_table_to_flowable(block, styles))
    return flow


def plain_text_flowables(text: str, styles: dict[str, ParagraphStyle]) -> list:
    flow: list = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            flow.append(Spacer(1, 2))
            continue
        if re.match(r"^#{1,6}\s+", stripped):
            flow.append(p(re.sub(r"^#{1,6}\s+", "", stripped), styles["h1"]))
        else:
            flow.append(p(stripped, styles["body"]))
    return flow


def build_document_pdf(
    source: SourceFile,
    output_pdf: Path,
    title: str,
    styles: dict[str, ParagraphStyle],
) -> None:
    text_dir = source.path.parent
    flow: list = [
        p(title, styles["title"]),
        p(source.rel, styles["meta"]),
        p(f"원본 형식: {source.ext.lstrip('.') or 'unknown'}", styles["meta"]),
        Spacer(1, 4),
    ]

    if source.ext == ".md":
        raw = source.path.read_text(encoding="utf-8", errors="ignore")
        flow.extend(markdown_flowables(raw, text_dir, styles))
    elif source.ext == ".html":
        raw = source.path.read_text(encoding="utf-8", errors="ignore")
        flow.extend(html_flowables(raw, text_dir, styles))
    elif source.ext == ".txt":
        raw = source.path.read_text(encoding="utf-8", errors="ignore")
        flow.extend(plain_text_flowables(raw, styles))
    elif source.ext == ".csv":
        flow.extend(csv_flowables(source.path, styles))
    elif source.ext == ".docx":
        flow.extend(docx_flowables(source.path, styles))
    elif source.ext in {".png", ".jpg", ".jpeg", ".webp"}:
        flow.extend(image_page(source.path, styles, caption=source.rel))
    elif source.ext == ".pdf":
        raise RuntimeError("PDF should be handled separately")
    else:
        raw = source.path.read_text(encoding="utf-8", errors="ignore")
        flow.extend(plain_text_flowables(raw, styles))

    doc = SimpleDocTemplate(
        str(output_pdf),
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=title,
        author=OWNER,
        subject=source.rel,
    )
    doc.build(flow, canvasmaker=NumberedCanvas)


def copy_pdf(source: Path, output_pdf: Path, title: str, styles: dict[str, ParagraphStyle]) -> None:
    reader = PdfReader(str(source))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.add_metadata({
        "/Title": title,
        "/Author": OWNER,
        "/Subject": source.name,
    })
    with output_pdf.open("wb") as f:
        writer.write(f)


def page_count(path: Path) -> int:
    return len(PdfReader(str(path)).pages)


def build_index(rows: list[dict[str, str]]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with INDEX_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["source", "title", "date", "output", "kind", "pages"])
        writer.writeheader()
        writer.writerows(rows)
    with INDEX_MD.open("w", encoding="utf-8") as f:
        f.write(f"# {PREFIX} index\n\n")
        f.write("| source | title | date | output | kind | pages |\n")
        f.write("| --- | --- | --- | --- | --- | ---: |\n")
        for row in rows:
            f.write(
                f"| {row['source']} | {row['title']} | {row['date']} | {row['output']} | {row['kind']} | {row['pages']} |\n"
            )


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    font_name = register_fonts()
    styles = make_styles(font_name)
    if OUT_DIR.exists():
        shutil.rmtree(OUT_DIR, ignore_errors=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    temp_root = extract_outer_inner_zip()
    try:
        sources = iter_source_files(temp_root)
        if not sources:
            raise RuntimeError("No files found under target prefix")

        source_index = build_source_index(sources)
        date_map: dict[str, str | None] = {source.rel: derive_created_date(source) for source in sources}
        for source in sources:
            if not date_map[source.rel]:
                date_map[source.rel] = find_parent_page_date(source, source_index, date_map, temp_root)
        for source in sources:
            if not date_map[source.rel]:
                date_map[source.rel] = DATE_STR

        used: dict[str, int] = {}
        rows: list[dict[str, str]] = []
        errors: list[str] = []
        for i, source in enumerate(sources, 1):
            try:
                title = derive_title(source)
                date_str = date_map.get(source.rel) or DATE_STR
                out_name = make_output_name(f"{title}_{date_str}", used)
                out_path = OUT_DIR / out_name
                if source.ext == ".pdf":
                    copy_pdf(source.path, out_path, title, styles)
                else:
                    build_document_pdf(source, out_path, title, styles)
                rows.append(
                    {
                        "source": source.rel,
                        "title": title,
                        "date": date_str,
                        "output": out_name,
                        "kind": source.ext.lstrip(".") or "unknown",
                        "pages": str(page_count(out_path)),
                    }
                )
                print(f"[{i}/{len(sources)}] {source.rel} -> {out_name}")
            except Exception as exc:
                msg = f"{source.rel} :: {type(exc).__name__}: {exc}"
                errors.append(msg)
                print(f"[{i}/{len(sources)}] ERROR {msg}")

        build_index(rows)
        if errors:
            err_path = OUT_DIR / f"{PREFIX}_errors.txt"
            err_path.write_text("\n".join(errors), encoding="utf-8")
            print(f"Errors: {len(errors)} -> {err_path}")
        print(f"PDF count: {len(rows)}")
        print(f"Output: {OUT_DIR}")
        print(f"Index: {INDEX_CSV}")
    finally:
        shutil.rmtree(temp_root, ignore_errors=True)


if __name__ == "__main__":
    main()
